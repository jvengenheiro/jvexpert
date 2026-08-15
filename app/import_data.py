"""Importação de instrutores e cursos a partir de arquivos CSV, Excel (.xlsx/.xls),
Word (.docx) ou PDF.

Pensado para receber exportações do SIGOP (ou de qualquer outro sistema) sem exigir um
formato de coluna exato: os cabeçalhos são normalizados e comparados contra uma lista de
sinônimos por campo (`CAMPO_SINONIMOS`), então "Nome do Curso", "curso" e "descricao"
mapeiam todos para o campo canônico `nome`, por exemplo. Com base nas colunas
reconhecidas, `detectar_tipo` tenta identificar sozinho se o arquivo é uma lista de
instrutores, cursos fixos ou cursos livres.

PDF e Word só são lidos quando contêm uma tabela de verdade (célula por célula) — texto
solto num PDF escaneado ou num relatório em prosa não é interpretado, pois qualquer
tentativa de "adivinhar" dados nesse cenário tende a errar silenciosamente.
"""

import re
import unicodedata
from datetime import datetime

import pandas as pd

from app.models import CursoFixo, CursoLivre, Instrutor, Qualificacao, db

DIAS_SEMANA_NOMES = {
    "segunda": 0,
    "segunda-feira": 0,
    "terca": 1,
    "terça": 1,
    "terca-feira": 1,
    "terça-feira": 1,
    "quarta": 2,
    "quarta-feira": 2,
    "quinta": 3,
    "quinta-feira": 3,
    "sexta": 4,
    "sexta-feira": 4,
    "sabado": 5,
    "sábado": 5,
    "domingo": 6,
}

# Sinônimos de cabeçalho por campo canônico. Comparados após normalização
# (sem acento, minúsculo, separadores viram "_"), então não precisam cobrir
# toda variação de maiúsculas/pontuação — só o "miolo" da palavra.
CAMPO_SINONIMOS = {
    "nome": [
        "nome", "curso", "nome_curso", "nome_do_curso", "descricao", "titulo", "curso_nome",
        "nome_do_instrutor", "nome_instrutor", "nome_completo",
    ],
    "categoria": ["categoria", "categoria_curso", "tipo", "tipo_curso", "area", "modalidade"],
    "carga_horaria": ["carga_horaria", "carga_horaria_total", "ch", "horas", "duracao", "carga"],
    "dia_semana": ["dia_semana", "dia_da_semana", "dia"],
    "hora_inicio": ["hora_inicio", "horario_inicio", "inicio", "hora_de_inicio", "h_inicio"],
    "hora_fim": ["hora_fim", "horario_fim", "fim", "termino", "hora_termino", "hora_de_termino", "h_fim"],
    "data_inicio": ["data_inicio", "data_de_inicio", "inicio_periodo", "data_inicial"],
    "data_fim": ["data_fim", "data_de_fim", "fim_periodo", "data_final", "data_termino"],
    "data": ["data", "data_curso", "data_da_turma", "data_turma", "data_realizacao"],
    "vagas": ["vagas", "qtd_vagas", "quantidade_vagas", "capacidade", "numero_vagas"],
    "instrutor": ["instrutor", "professor", "responsavel", "instrutor_titular", "docente"],
    "email": ["email", "e_mail", "correio_eletronico"],
    "telefone": ["telefone", "celular", "contato", "fone", "whatsapp"],
    "categorias": ["categorias", "qualificacoes", "habilitacoes", "especialidades"],
    "carga_horaria_semanal_max": [
        "carga_horaria_semanal_max", "carga_horaria_semanal", "horas_semana", "limite_semanal",
    ],
    "sigop_id": ["sigop_id", "id_sigop", "codigo_sigop", "codigo", "id"],
}

CAMPOS_OBRIGATORIOS = {
    "instrutores": ["nome"],
    "cursos_fixos": ["nome", "categoria", "dia_semana", "hora_inicio", "hora_fim", "data_inicio", "data_fim"],
    "cursos_livres": ["nome", "categoria", "data", "hora_inicio", "hora_fim"],
}

FORMATOS_SUPORTADOS = (".csv", ".xlsx", ".xls", ".docx", ".pdf")


def _normalizar(texto):
    texto = unicodedata.normalize("NFKD", str(texto or "")).encode("ascii", "ignore").decode("ascii")
    texto = texto.strip().lower()
    return re.sub(r"[^a-z0-9]+", "_", texto).strip("_")


def _mapear_colunas(colunas, campos_relevantes):
    """Para cada campo canônico relevante, encontra a coluna original correspondente
    (se houver) comparando versões normalizadas contra a lista de sinônimos."""
    sinonimo_para_campo = {}
    for campo in campos_relevantes:
        for sinonimo in CAMPO_SINONIMOS.get(campo, [campo]):
            sinonimo_para_campo[_normalizar(sinonimo)] = campo

    mapa = {}
    for coluna in colunas:
        campo = sinonimo_para_campo.get(_normalizar(coluna))
        if campo and campo not in mapa:
            mapa[campo] = coluna
    return mapa


def _renomear_linhas(rows, mapa):
    """Reescreve as chaves de cada linha usando os nomes de campo canônicos, com base
    no mapa {campo_canonico: coluna_original} descoberto por `_mapear_colunas`."""
    coluna_para_campo = {coluna: campo for campo, coluna in mapa.items()}
    linhas_renomeadas = []
    for linha in rows:
        nova = {}
        for chave, valor in linha.items():
            nova[coluna_para_campo.get(chave, chave)] = valor
        linhas_renomeadas.append(nova)
    return linhas_renomeadas


def _linhas_de_celulas_para_registros(linhas_de_celulas):
    """Converte uma lista de linhas (cada uma uma lista de células de texto) — vinda de
    uma tabela de PDF/Word — em (registros, colunas), usando a primeira linha não vazia
    como cabeçalho."""
    linhas = [
        [(c or "").strip() for c in linha]
        for linha in linhas_de_celulas
        if any((c or "").strip() for c in linha)
    ]
    if not linhas:
        raise ValueError("tabela vazia")

    cabecalho = linhas[0]
    registros = []
    for linha in linhas[1:]:
        if linha == cabecalho:
            continue  # cabeçalho repetido em páginas seguintes de um PDF
        registro = {}
        for idx, valor in enumerate(linha):
            if idx < len(cabecalho) and cabecalho[idx]:
                registro[cabecalho[idx]] = valor
        registros.append(registro)
    return registros, cabecalho


def _ler_docx(arquivo):
    from docx import Document

    documento = Document(arquivo)
    if not documento.tables:
        raise ValueError(
            "nenhuma tabela encontrada no documento Word (.docx) — o arquivo precisa ter uma "
            "tabela de verdade, não apenas texto corrido"
        )
    tabela = documento.tables[0]
    linhas_de_celulas = [[celula.text for celula in linha.cells] for linha in tabela.rows]
    return _linhas_de_celulas_para_registros(linhas_de_celulas)


def _ler_pdf(arquivo):
    import pdfplumber

    linhas_de_celulas = []
    with pdfplumber.open(arquivo) as pdf:
        for pagina in pdf.pages:
            for tabela in pagina.extract_tables() or []:
                linhas_de_celulas.extend(tabela)

    if not linhas_de_celulas:
        raise ValueError(
            "nenhuma tabela encontrada no PDF — funciona melhor com PDFs gerados a partir de "
            "planilhas/relatórios tabulares (ex: exportação direta do SIGOP), não com "
            "documentos escaneados ou em texto corrido"
        )
    return _linhas_de_celulas_para_registros(linhas_de_celulas)


def ler_tabela(arquivo):
    """Lê qualquer um dos formatos suportados e retorna (linhas, colunas), onde `linhas`
    é uma lista de dicts {coluna_original: valor} e `colunas` é a lista de cabeçalhos na
    ordem original — antes de qualquer mapeamento para os campos canônicos."""
    nome = (arquivo.filename or "").lower()

    if nome.endswith(".xlsx") or nome.endswith(".xls"):
        df = pd.read_excel(arquivo)
    elif nome.endswith(".csv"):
        df = pd.read_csv(arquivo)
    elif nome.endswith(".docx"):
        return _ler_docx(arquivo)
    elif nome.endswith(".pdf"):
        return _ler_pdf(arquivo)
    else:
        raise ValueError(
            f"formato de arquivo não suportado: {nome or 'desconhecido'} "
            f"(use um dos formatos: {', '.join(FORMATOS_SUPORTADOS)})"
        )

    df = df.fillna("")
    colunas = [str(c) for c in df.columns]
    linhas = [{str(k): v for k, v in registro.items()} for registro in df.to_dict("records")]
    return linhas, colunas


def detectar_tipo(colunas):
    """Tenta identificar sozinho se as colunas descrevem instrutores, cursos fixos ou
    cursos livres. Retorna None quando ambíguo — nesse caso o usuário escolhe manualmente."""
    mapa = _mapear_colunas(colunas, list(CAMPO_SINONIMOS.keys()))
    campos = set(mapa.keys())

    if "dia_semana" in campos:
        return "cursos_fixos"
    if "data" in campos and "categoria" in campos:
        return "cursos_livres"
    if "categorias" in campos or "email" in campos or "telefone" in campos:
        return "instrutores"
    if "nome" in campos and "categoria" not in campos and "data" not in campos:
        return "instrutores"
    return None


def _preparar_linhas(rows, colunas, tipo):
    mapa = _mapear_colunas(colunas, CAMPOS_OBRIGATORIOS[tipo] + ["carga_horaria", "vagas", "instrutor", "sigop_id"])
    faltando = [campo for campo in CAMPOS_OBRIGATORIOS[tipo] if campo not in mapa]
    if faltando:
        raise ValueError(
            "não foi possível identificar as colunas: "
            + ", ".join(faltando)
            + ". Colunas encontradas no arquivo: "
            + ", ".join(colunas)
        )
    return _renomear_linhas(rows, mapa)


def _parse_dia_semana(valor):
    texto = str(valor).strip().lower()
    if texto in DIAS_SEMANA_NOMES:
        return DIAS_SEMANA_NOMES[texto]
    try:
        dia = int(float(texto))
    except ValueError as exc:
        raise ValueError(f"dia_semana inválido: {valor!r}") from exc
    if not 0 <= dia <= 6:
        raise ValueError(f"dia_semana fora do intervalo 0-6: {valor!r}")
    return dia


def _parse_hora(valor):
    texto = str(valor).strip()
    for formato in ("%H:%M", "%H:%M:%S"):
        try:
            return datetime.strptime(texto, formato).time()
        except ValueError:
            continue
    raise ValueError(f"horário inválido (use HH:MM): {valor!r}")


def _parse_data(valor):
    if isinstance(valor, datetime):
        return valor.date()
    texto = str(valor).strip()
    for formato in ("%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(texto, formato).date()
        except ValueError:
            continue
    raise ValueError(f"data inválida (use AAAA-MM-DD): {valor!r}")


def _linha_para_erro(indice, exc):
    return f"linha {indice + 2}: {exc}"


def _valor_opcional(linha, coluna):
    if coluna not in linha or linha[coluna] is None or pd.isna(linha[coluna]):
        return None
    valor = str(linha[coluna]).strip()
    return valor or None


def importar_instrutores(arquivo):
    rows, colunas = ler_tabela(arquivo)
    return processar_instrutores(rows, colunas)


def importar_cursos_fixos(arquivo):
    rows, colunas = ler_tabela(arquivo)
    return processar_cursos_fixos(rows, colunas)


def importar_cursos_livres(arquivo):
    rows, colunas = ler_tabela(arquivo)
    return processar_cursos_livres(rows, colunas)


def processar_instrutores(rows, colunas):
    rows = _preparar_linhas(rows, colunas, "instrutores")
    erros = []
    criados = 0
    for indice, linha in enumerate(rows):
        try:
            nome = str(linha.get("nome", "")).strip()
            if not nome or nome.lower() == "nan":
                raise ValueError("nome é obrigatório")

            instrutor = Instrutor.query.filter_by(nome=nome).first()
            if instrutor is None:
                instrutor = Instrutor(nome=nome)
                db.session.add(instrutor)

            instrutor.email = _valor_opcional(linha, "email")
            instrutor.telefone = _valor_opcional(linha, "telefone")
            instrutor.sigop_id = _valor_opcional(linha, "sigop_id")

            carga = _valor_opcional(linha, "carga_horaria_semanal_max")
            instrutor.carga_horaria_semanal_max = int(float(carga)) if carga else None

            categorias_raw = _valor_opcional(linha, "categorias")
            if categorias_raw:
                instrutor.qualificacoes = [
                    Qualificacao(categoria=c.strip())
                    for c in categorias_raw.split(";")
                    if c.strip()
                ]

            criados += 1
        except Exception as exc:  # noqa: BLE001 - reportamos linha a linha
            erros.append(_linha_para_erro(indice, exc))

    db.session.commit()
    return criados, erros


def processar_cursos_fixos(rows, colunas):
    rows = _preparar_linhas(rows, colunas, "cursos_fixos")
    erros = []
    criados = 0
    for indice, linha in enumerate(rows):
        try:
            instrutor_id = None
            instrutor_nome = _valor_opcional(linha, "instrutor")
            if instrutor_nome:
                instrutor = Instrutor.query.filter_by(nome=instrutor_nome).first()
                if instrutor is None:
                    raise ValueError(f"instrutor não encontrado: {instrutor_nome!r}")
                instrutor_id = instrutor.id

            carga = _valor_opcional(linha, "carga_horaria")
            vagas = _valor_opcional(linha, "vagas")

            curso = CursoFixo(
                nome=str(linha["nome"]).strip(),
                categoria=str(linha["categoria"]).strip(),
                carga_horaria=int(float(carga)) if carga else None,
                dia_semana=_parse_dia_semana(linha["dia_semana"]),
                hora_inicio=_parse_hora(linha["hora_inicio"]),
                hora_fim=_parse_hora(linha["hora_fim"]),
                data_inicio=_parse_data(linha["data_inicio"]),
                data_fim=_parse_data(linha["data_fim"]),
                vagas=int(float(vagas)) if vagas else None,
                instrutor_fixo_id=instrutor_id,
                sigop_id=_valor_opcional(linha, "sigop_id"),
            )
            db.session.add(curso)
            criados += 1
        except Exception as exc:  # noqa: BLE001
            erros.append(_linha_para_erro(indice, exc))

    db.session.commit()
    return criados, erros


def processar_cursos_livres(rows, colunas):
    rows = _preparar_linhas(rows, colunas, "cursos_livres")
    erros = []
    criados = 0
    for indice, linha in enumerate(rows):
        try:
            carga = _valor_opcional(linha, "carga_horaria")
            vagas = _valor_opcional(linha, "vagas")

            curso = CursoLivre(
                nome=str(linha["nome"]).strip(),
                categoria=str(linha["categoria"]).strip(),
                carga_horaria=int(float(carga)) if carga else None,
                data=_parse_data(linha["data"]),
                hora_inicio=_parse_hora(linha["hora_inicio"]),
                hora_fim=_parse_hora(linha["hora_fim"]),
                vagas=int(float(vagas)) if vagas else None,
                sigop_id=_valor_opcional(linha, "sigop_id"),
            )
            db.session.add(curso)
            criados += 1
        except Exception as exc:  # noqa: BLE001
            erros.append(_linha_para_erro(indice, exc))

    db.session.commit()
    return criados, erros
